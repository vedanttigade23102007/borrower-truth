"""Read a borrower's spreadsheet as untrusted input: values only, never formulas, nothing that can run.

- CSV, or Excel .xlsx. A macro workbook (.xlsm, or an .xlsx that carries a macro project) is refused.
- Files over MAX_BYTES are refused, and so is an .xlsx whose contents would unpack to more than MAX_UNPACKED
  (a "zip bomb": a small file that expands to fill the disk).
- Excel cells are read as their stored values (`data_only=True`): a formula is never read as text and never
  evaluated. A formula whose value was never saved reads as empty, and the checks then reject that row.
- XML is parsed through defusedxml (openpyxl uses it when installed), so entity-expansion and external-entity
  tricks in the workbook's XML are refused.
"""
from __future__ import annotations

import csv
import os
import zipfile

MAX_BYTES = 50 * 1024 * 1024
MAX_UNPACKED = 500 * 1024 * 1024
MAX_ROWS = 2_000_000


class Refused(Exception):
    """The file is not read at all."""


def kind_of(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsm", ".xlsb", ".xls", ".xltm"):
        raise Refused(f"{ext} files are refused: macro or legacy workbooks. Ask for .xlsx or .csv")
    if ext == ".xlsx":
        return "xlsx"
    if ext in (".csv", ".txt"):
        return "csv"
    raise Refused(f"unknown file type '{ext}': only .xlsx and .csv are read")


def _guard_size(path: str, max_bytes: int) -> None:
    size = os.path.getsize(path)
    if size > max_bytes:
        raise Refused(f"file is {size:,} bytes, over the {max_bytes:,} byte limit")


def _guard_xlsx(path: str, max_unpacked: int) -> None:
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise Refused("not a valid .xlsx file")
    with z:
        names = [i.filename.lower() for i in z.infolist()]
        if any(n.endswith("vbaproject.bin") for n in names):
            raise Refused("the workbook carries a macro project (renamed .xlsm?): refused")
        types = z.read("[Content_Types].xml").decode("utf-8", "replace").lower() if "[content_types].xml" in names else ""
        if "macroenabled" in types or "vbaproject" in types:
            raise Refused("the workbook is macro-enabled: refused")
        unpacked = sum(i.file_size for i in z.infolist())
        if unpacked > max_unpacked:
            raise Refused(f"the workbook unpacks to {unpacked:,} bytes, over the {max_unpacked:,} byte limit")


def sheets(path: str, max_bytes: int = MAX_BYTES, max_unpacked: int = MAX_UNPACKED) -> list[str]:
    kind = kind_of(path)
    _guard_size(path, max_bytes)
    if kind == "csv":
        return ["(csv)"]
    _guard_xlsx(path, max_unpacked)
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def rows(path: str, sheet: str | None = None, max_bytes: int = MAX_BYTES, max_unpacked: int = MAX_UNPACKED,
         max_rows: int = MAX_ROWS, delimiter: str = ",") -> list[list]:
    """Every row of one sheet, as values (row 1 is the first). CSV values stay text; Excel gives numbers, text,
    dates and booleans as stored."""
    kind = kind_of(path)
    _guard_size(path, max_bytes)
    out: list[list] = []
    if kind == "csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            for r in csv.reader(f, delimiter=delimiter):
                out.append(r)
                if len(out) > max_rows:
                    raise Refused(f"more than {max_rows:,} rows")
        return out
    _guard_xlsx(path, max_unpacked)
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet is None:
            sheet = wb.sheetnames[0]
        if sheet not in wb.sheetnames:
            raise Refused(f"no sheet named '{sheet}' (sheets: {', '.join(wb.sheetnames)})")
        for r in wb[sheet].iter_rows(values_only=True):
            out.append(list(r))
            if len(out) > max_rows:
                raise Refused(f"more than {max_rows:,} rows")
    finally:
        wb.close()
    return out
